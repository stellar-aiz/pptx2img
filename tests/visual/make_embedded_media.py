"""Independent fixture: a slide with an embedded video larger than the 32 MiB part limit.

The video payload is generated zero bytes (never rendered) and the poster frame is a generated
PNG. The deck is about 40 MiB uncompressed, so write it outside the repository rather than
committing it. python-pptx and Pillow are needed only for generation.
"""
from io import BytesIO
from pathlib import Path
import sys
import tempfile

from PIL import Image
from pptx import Presentation
from pptx.util import Inches

MIB = 1024 * 1024


def main(out: Path, size_mib: int = 40):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(10), Inches(5.625)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    poster = BytesIO()
    Image.new("RGB", (160, 90), (0, 102, 204)).save(poster, "PNG")
    poster.seek(0)
    with tempfile.TemporaryDirectory() as tmp:
        video = Path(tmp) / "dummy.mov"
        video.write_bytes(bytes(size_mib * MIB))
        slide.shapes.add_movie(str(video), Inches(1), Inches(1), Inches(8), Inches(3.5),
                               poster_frame_image=poster, mime_type="video/quicktime")
        prs.save(out)


if __name__ == "__main__":
    main(Path(sys.argv[1]), *(int(arg) for arg in sys.argv[2:3]))
