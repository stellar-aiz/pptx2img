"""Independent fixture for text inside ellipse, roundRect, and rect presets.

python-pptx is needed only for generation.

One slide, all labels use the same font and sizes in both renderers:

- Left: three concentric circles: the outer two have top-anchored labels (each must stay
  inside its circle's inscribed rectangle, not start at the corner of its box) and the
  inner one has a centred label.
- Right top: two overlapping ellipses with Japanese text, one top-anchored and one centred.
- Right bottom: a roundRect with `adj` 50000, a plain rect, and a small ellipse whose
  text is taller than its inscribed rectangle (records `text_overflow`).
"""
import argparse
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt


def shape(slide, kind, x, y, w, h, fill, anchor, text, size, align=PP_ALIGN.CENTER, adj=None):
    item = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    style = item._element.find(qn("p:style"))
    if style is not None:
        item._element.remove(style)
    if adj is not None:
        item.adjustments[0] = adj
    item.fill.solid()
    item.fill.fore_color.rgb = RGBColor(*fill)
    item.line.color.rgb = RGBColor(40, 60, 90)
    item.line.width = Pt(1)
    frame = item.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = Pt(0)
    frame.vertical_anchor = anchor
    for index, line in enumerate(text.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = align
        run = paragraph.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.name = "Arial"
        run.font.color.rgb = RGBColor(0, 0, 0)
        rpr = run._r.get_or_add_rPr()
        east_asian = rpr.makeelement(qn("a:ea"), {"typeface": "Hiragino Sans"})
        rpr.find(qn("a:latin")).addnext(east_asian)
    return item


def make(output):
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(10), Inches(5.6)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    oval = MSO_SHAPE.OVAL
    # Concentric circles share one centre; each label sits at the top of its own ring.
    shape(slide, oval, 0.3, 0.3, 5.0, 5.0, (214, 228, 247), MSO_ANCHOR.TOP, "外側の円", 10)
    shape(slide, oval, 1.3, 1.3, 3.0, 3.0, (255, 233, 200), MSO_ANCHOR.TOP, "中間の円", 10)
    shape(slide, oval, 2.0, 2.0, 1.6, 1.6, (210, 238, 214), MSO_ANCHOR.MIDDLE, "内側の円の中心に置く日本語ラベル", 11)
    # Overlapping ellipses: the upper one covers the lower one's box corner, not its text.
    shape(slide, oval, 5.6, 0.3, 2.6, 1.9, (252, 220, 220), MSO_ANCHOR.TOP, "楕円 A の上寄せ日本語テキストは内接矩形に収まる", 11)
    shape(slide, oval, 7.0, 1.1, 2.6, 1.9, (225, 220, 250), MSO_ANCHOR.MIDDLE, "楕円 B の中央寄せ日本語テキストは内接矩形に収まる", 11)
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 5.6, 3.4, 2.2, 0.9, (240, 240, 200), MSO_ANCHOR.MIDDLE, "角丸の長い日本語ラベルが半径分だけ内側に収まる", 10, adj=0.5)
    shape(slide, MSO_SHAPE.RECTANGLE, 8.0, 3.4, 1.6, 0.9, (230, 230, 230), MSO_ANCHOR.TOP, "矩形は全面を使う長い日本語ラベル", 10, align=PP_ALIGN.LEFT)
    shape(slide, oval, 6.5, 4.4, 1.0, 1.0, (255, 210, 210), MSO_ANCHOR.TOP, "小さな楕円に収まらない長い日本語の文章をさらに続けて書く", 11)
    deck.save(output)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("output", type=Path)
    make(parser.parse_args().output)
