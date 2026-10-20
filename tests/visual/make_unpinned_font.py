"""Independent fixture for the unpinned-family report: text that requests a family which no
substitution table lists. python-pptx is needed only for generation."""
from pathlib import Path
import sys

from pptx import Presentation
from pptx.util import Inches, Pt

# Not a real font; the name must stay absent from tests/visual/font-substitution*.json.
UNLISTED_FAMILY = "Pptx2img Unlisted Sans"
TEXT = "Unlisted family sample 0123"


def main(out: Path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(10), Inches(5.625)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(0.8))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = f"{TEXT} ({UNLISTED_FAMILY})"
    run.font.name = UNLISTED_FAMILY
    run.font.size = Pt(24)
    prs.save(out)


if __name__ == "__main__":
    main(Path(sys.argv[1]))
