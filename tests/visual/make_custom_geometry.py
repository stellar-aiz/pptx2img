"""Independent fixture for custom geometry (a:custGeom) freeforms; no source-document assets.

python-pptx is needed only for generation. Six cells, no text (font differences would
blur the outline comparison):

- top row: a closed polygon from python-pptx's freeform builder (moveTo/lnTo/close),
  a panel whose lower edge is a cubic Bezier, and a tab with arcTo corners;
- bottom row: an open quadratic curve with fill="none" on a shape that has a solid
  fill, a bracket built from gdLst guides with a path in shape coordinates (no path
  w/h), and a stroke="0" disk of two arcs under a separate fill="none" square outline path.

Next to the deck, `<out>.json` lists one comparison region per cell for compare.py
`--regions`, so a missing freeform is caught even when the whole page passes.
"""
import json
from pathlib import Path
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml import parse_xml
from pptx.util import Inches, Pt

A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
P = 'http://schemas.openxmlformats.org/presentationml/2006/main'
SLIDE_W, SLIDE_H = Inches(10), Inches(5.625)
W, H = Inches(2.8), Inches(2.2)
COLUMNS = [Inches(0.4), Inches(3.6), Inches(6.8)]
ROWS = [Inches(0.4), Inches(3.0)]


def pt(x, y):
    return f'<a:pt x="{x}" y="{y}"/>'


def style(shape, fill, line='202020', width=Pt(3)):
    """Solid fill (or none) and a solid outline, without a theme style reference."""
    element = shape._element
    for tag in ('style', 'txBody'):
        found = element.find(f'{{{P}}}{tag}')
        if found is not None:
            element.remove(found)
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(fill)
    else:
        shape.fill.background()
    shape.line.color.rgb = RGBColor.from_string(line)
    shape.line.width = width


def freeform(slide, x, y, geometry, fill, **line):
    """A rectangle whose preset geometry is replaced by the given a:custGeom body."""
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, W, H)
    spPr = shape._element.spPr
    preset = spPr.find(f'{{{A}}}prstGeom')
    preset.addnext(parse_xml(f'<a:custGeom xmlns:a="{A}">{geometry}</a:custGeom>'))
    spPr.remove(preset)
    style(shape, fill, **line)
    return shape


def main(out):
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # 1. python-pptx freeform builder: a closed polygon of straight segments.
    x, y = COLUMNS[0], ROWS[0]
    builder = slide.shapes.build_freeform(x + W // 2, y)
    builder.add_line_segments(
        [(x + W, y + H * 2 // 5), (x + W * 4 // 5, y + H), (x + W // 5, y + H), (x, y + H * 2 // 5)],
        close=True,
    )
    style(builder.convert_to_shape(), '4FA394')

    # 2. Lines plus a cubic Bezier lower edge in a 1000x600 path space.
    freeform(slide, COLUMNS[1], ROWS[0], (
        '<a:avLst/><a:gdLst/><a:rect l="l" t="t" r="r" b="b"/><a:pathLst>'
        '<a:path w="1000" h="600">'
        f'<a:moveTo>{pt(0, 0)}</a:moveTo><a:lnTo>{pt(1000, 0)}</a:lnTo><a:lnTo>{pt(1000, 450)}</a:lnTo>'
        f'<a:cubicBezTo>{pt(700, 650)}{pt(300, 250)}{pt(0, 600)}</a:cubicBezTo><a:close/>'
        '</a:path></a:pathLst>'
    ), 'D2EDFE', line='245A51')

    # 3. A tab with quarter-ellipse top corners drawn by arcTo (path in shape EMU).
    w, h, r = int(W), int(H), int(Inches(0.6))
    freeform(slide, COLUMNS[2], ROWS[0], (
        '<a:pathLst>'
        f'<a:path w="{w}" h="{h}">'
        f'<a:moveTo>{pt(0, h)}</a:moveTo><a:lnTo>{pt(0, r)}</a:lnTo>'
        f'<a:arcTo wR="{r}" hR="{r}" stAng="10800000" swAng="5400000"/>'
        f'<a:lnTo>{pt(w - r, 0)}</a:lnTo>'
        f'<a:arcTo wR="{r}" hR="{r}" stAng="16200000" swAng="5400000"/>'
        f'<a:lnTo>{pt(w, h)}</a:lnTo><a:close/>'
        '</a:path></a:pathLst>'
    ), 'FFC000')

    # 4. An open quadratic curve: fill="none" on the path keeps the shape fill out.
    freeform(slide, COLUMNS[0], ROWS[1], (
        '<a:pathLst><a:path w="100" h="100" fill="none">'
        f'<a:moveTo>{pt(0, 100)}</a:moveTo><a:quadBezTo>{pt(25, 0)}{pt(50, 50)}</a:quadBezTo>'
        f'<a:quadBezTo>{pt(75, 100)}{pt(100, 0)}</a:quadBezTo>'
        '</a:path></a:pathLst>'
    ), 'C00000', line='0070C0', width=Pt(6))

    # 5. A bracket from gdLst guides; the path has no w/h, so it uses shape EMU.
    freeform(slide, COLUMNS[1], ROWS[1], (
        '<a:avLst/><a:gdLst>'
        '<a:gd name="arm" fmla="*/ w 1 4"/>'
        '<a:gd name="bar" fmla="*/ ss 1 6"/>'
        '<a:gd name="inner" fmla="+- r 0 bar"/>'
        '<a:gd name="low" fmla="+- b 0 bar"/>'
        '<a:gd name="armIn" fmla="+- r 0 arm"/>'
        '</a:gdLst><a:rect l="l" t="bar" r="inner" b="low"/><a:pathLst><a:path>'
        f'<a:moveTo>{pt("armIn", "t")}</a:moveTo><a:lnTo>{pt("r", "t")}</a:lnTo>'
        f'<a:lnTo>{pt("r", "b")}</a:lnTo><a:lnTo>{pt("armIn", "b")}</a:lnTo>'
        f'<a:lnTo>{pt("armIn", "low")}</a:lnTo><a:lnTo>{pt("inner", "low")}</a:lnTo>'
        f'<a:lnTo>{pt("inner", "bar")}</a:lnTo><a:lnTo>{pt("armIn", "bar")}</a:lnTo><a:close/>'
        '</a:path></a:pathLst>'
    ), '7030A0')

    # 6. A stroke="0" disk filled by two half-turn arcs, outlined by a separate fill="none" path.
    disk = (
        f'<a:moveTo>{pt(0, 50)}</a:moveTo>'
        '<a:arcTo wR="50" hR="50" stAng="10800000" swAng="10800000"/>'
        '<a:arcTo wR="50" hR="50" stAng="0" swAng="10800000"/><a:close/>'
    )
    freeform(slide, COLUMNS[2], ROWS[1], (
        '<a:pathLst>'
        f'<a:path w="100" h="100" stroke="0">{disk}</a:path>'
        f'<a:path w="100" h="100" fill="none"><a:moveTo>{pt(0, 0)}</a:moveTo>'
        f'<a:lnTo>{pt(100, 0)}</a:lnTo><a:lnTo>{pt(100, 100)}</a:lnTo><a:lnTo>{pt(0, 100)}</a:lnTo><a:close/></a:path>'
        '</a:pathLst>'
    ), '00B0F0', line='C00000')

    prs.save(out)
    regions = [
        {
            'slide': 1,
            'name': f'cell-{row + 1}-{column + 1}',
            'rect': [
                round((int(COLUMNS[column]) - 40000) / SLIDE_W, 4),
                round((int(ROWS[row]) - 40000) / SLIDE_H, 4),
                round((int(W) + 80000) / SLIDE_W, 4),
                round((int(H) + 80000) / SLIDE_H, 4),
            ],
        }
        for row in range(2)
        for column in range(3)
    ]
    Path(out).with_suffix('.json').write_text(json.dumps({'regions': regions}, indent=2) + '\n')


if __name__ == '__main__':
    main(Path(sys.argv[1]))
