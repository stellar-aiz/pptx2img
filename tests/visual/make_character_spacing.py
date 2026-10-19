"""Character-spacing comparison fixtures for this project; python-pptx is needed only for regeneration.

spc uses 1/100pt. Explicit negative values on punctuation runs are distinct from automatic compression.
"""
import argparse
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement


def make(output: Path):
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(6), Inches(3.5)

    def box(slide, text, y, *, mode="punct", align=PP_ALIGN.LEFT,
            width=5.6, x=.2, inherited=None, explicit=None, size=24):
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(1.5))
        frame = shape.text_frame
        frame.auto_size = MSO_AUTO_SIZE.NONE
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = Pt(0)
        frame.margin_top = frame.margin_bottom = Pt(0)
        para = frame.paragraphs[0]
        para.alignment = align
        para.line_spacing = Pt(30)
        props = para._p.get_or_add_pPr()
        props.set("hangingPunct", "0")
        if inherited is not None:
            props.get_or_add_defRPr().set("spc", str(inherited))
        for segment in list(text) if mode == "punct" else [text]:
            run = para.add_run()
            run.text = segment
            run.font.name, run.font.size = "Hiragino Sans", Pt(size)
            rp = run._r.get_or_add_rPr()
            rp.set("lang", "ja-JP")
            spacing = explicit if mode == "inherit" else (
                -1200 if mode == "all" or mode == "punct" and segment in "、。）」" else 0)
            if spacing is not None:
                rp.set("spc", str(spacing))
            ea = OxmlElement("a:ea")
            ea.set("typeface", "Hiragino Sans")
            rp.append(ea)

    text = "あ」、い。。う、、え）」お"
    for mode in ("punct", "all", "none"):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        for i, align in enumerate((PP_ALIGN.LEFT, PP_ALIGN.CENTER, PP_ALIGN.RIGHT)):
            box(slide, text, .1 + i * .8, mode=mode, align=align)
        # Verify that preceding character spacing does not leak across style boundaries.
        box(slide, "あ、い。う、え。お", 2.5, mode=mode)

    slide = deck.slides.add_slide(deck.slide_layouts[6])
    box(slide, "あ」、い。。う、、え）」お", .1, width=2.5)
    box(slide, "あ」、い。。う、、え）」お", .1, width=2.5, x=3.2, mode="none")
    box(slide, "あ、い。う、え。お", 1.9, width=2.5, mode="all")
    box(slide, "あいうえおかきく", 1.9, width=2.5, x=3.2, mode="inherit", explicit=300)

    slide = deck.slides.add_slide(deck.slide_layouts[6])
    # Paragraph default inheritance, explicit zero overrides, positive spacing, and different sizes.
    for i, (inherited, explicit, size) in enumerate(((-600,None,24),(-600,0,24),(-600,300,24),(-600,None,18))):
        box(slide, "あ」、い。。う、、え", .1 + i * .8, mode="inherit",
            inherited=inherited, explicit=explicit, size=size, align=PP_ALIGN.RIGHT)
    deck.save(output)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("output", type=Path)
    make(parser.parse_args().output)
