"""Independent fixture for slide-number fields in layout text boxes and in slide text.

Three slides share one layout whose footer text box holds a `slidenum` field saved as the
placeholder text `‹#›`, which is how PowerPoint stores it. Each slide also has a paragraph
with a run followed by a `slidenum` field, so the field text takes part in line wrapping.
python-pptx is needed only for generation.
"""
import argparse
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
ID = "{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}"


def field(kind="slidenum", saved="‹#›", size=1800):
    return (
        f'<a:fld xmlns:a="{A}" id="{ID}" type="{kind}">'
        f'<a:rPr lang="en-US" sz="{size}"><a:latin typeface="Arial"/></a:rPr><a:t>{saved}</a:t></a:fld>'
    )


def footer(layout):
    """Adds a plain (non-placeholder) text box with the field to the layout."""
    xml = (
        f'<p:sp xmlns:p="{P}" xmlns:a="{A}"><p:nvSpPr><p:cNvPr id="90" name="Footer number"/>'
        '<p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr><p:spPr><a:xfrm>'
        f'<a:off x="{Inches(5.2)}" y="{Inches(3.0)}"/><a:ext cx="{Inches(0.6)}" cy="{Inches(0.4)}"/></a:xfrm>'
        '<a:prstGeom prst="rect"/></p:spPr><p:txBody><a:bodyPr wrap="square" lIns="0" tIns="0" rIns="0" bIns="0"/>'
        f'<a:lstStyle/><a:p><a:pPr algn="r"/>{field()}</a:p></p:txBody></p:sp>'
    )
    layout.shapes._spTree.append(etree.fromstring(xml))


def make(output, first_slide_num=None):
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(6), Inches(3.5)
    if first_slide_num is not None:
        deck._element.set("firstSlideNum", str(first_slide_num))
    layout = deck.slide_layouts[6]
    footer(layout)
    for index in range(3):
        slide = deck.slides.add_slide(layout)
        box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(3.2), Inches(1.5))
        frame = box.text_frame
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        paragraph = frame.paragraphs[0]
        run = paragraph.add_run()
        run.text = f"Report section {index + 1} continues on page"
        run.font.size = Pt(18)
        run.font.name = "Arial"
        paragraph._p.append(etree.fromstring(field()))
    deck.save(output)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("output", type=Path)
    parser.add_argument("--first-slide-num", type=int, help="value for p:presentation/@firstSlideNum")
    arguments = parser.parse_args()
    make(arguments.output, arguments.first_slide_num)
