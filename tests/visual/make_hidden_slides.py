"""Independent fixture for hidden slides (p:sld/@show="0"). python-pptx is needed only for generation."""
from pathlib import Path
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

# (fill color, hidden)
SLIDES = [("2E75B6", False), ("C00000", True), ("548235", False), ("C00000", True), ("7030A0", False)]


def main(out: Path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(10), Inches(5.625)
    for number, (color, hidden) in enumerate(SLIDES, start=1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5 * number), Inches(1), Inches(2), Inches(2.5))
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor.from_string(color)
        box.line.fill.background()
        if hidden:
            slide._element.set("show", "0")
    prs.save(out)


if __name__ == "__main__":
    main(Path(sys.argv[1]))
