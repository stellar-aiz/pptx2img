"""階層付き記号、開始番号と連番、buNoneを検証する自作資料。"""
import argparse
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn


def make(output):
    deck=Presentation();deck.slide_width,deck.slide_height=Inches(6),Inches(3.5)
    for numbered in [False,True]:
        slide=deck.slides.add_slide(deck.slide_layouts[6])
        shape=slide.shapes.add_textbox(Inches(.3),Inches(.2),Inches(5.4),Inches(3.0))
        frame=shape.text_frame;frame.auto_size=MSO_AUTO_SIZE.NONE;frame.word_wrap=True
        frame.margin_top=frame.margin_bottom=frame.margin_left=frame.margin_right=Pt(0)
        listing=frame._txBody.find(qn('a:lstStyle'))
        for level in [0,1]:
            pp=OxmlElement(f'a:lvl{level+1}pPr');pp.set('marL',str((level+1)*228600));pp.set('indent','-228600')
            marker=OxmlElement('a:buAutoNum' if numbered else 'a:buChar')
            marker.set('type','arabicPeriod') if numbered else marker.set('char','•');pp.append(marker)
            font=OxmlElement('a:buFont');font.set('typeface','Arial');pp.append(font)
            clr=OxmlElement('a:buClr');rgb=OxmlElement('a:srgbClr');rgb.set('val','3366CC');clr.append(rgb);pp.append(clr)
            listing.append(pp)
        for index,level in enumerate([0,1,0,0]):
            para=frame.paragraphs[0] if index==0 else frame.add_paragraph()
            para.level=level;para.line_spacing=1.0;para.space_after=Pt(6)
            pp=para._p.get_or_add_pPr()
            if numbered and index==0:
                n=OxmlElement('a:buAutoNum');n.set('type','arabicPeriod');n.set('startAt','3');pp.append(n)
            if index==3:
                pp.append(OxmlElement('a:buNone'));pp.set('marL','0');pp.set('indent','0')
            run=para.add_run();run.text=('No bullet here.' if index==3 else 'Nested item with enough independent example text to wrap onto another line.')
            run.font.name='Arial';run.font.size=Pt(16)
    deck.save(output)

if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__);parser.add_argument('output',type=Path)
    make(parser.parse_args().output)
