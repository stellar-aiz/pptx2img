"""Independent fixture: an `a:tile` picture fill on a shape (a right triangle).

The picture is a 64x64 px checkerboard drawn with Pillow and stored without `pHYs`
density, so the tile size comes from the renderer's default density. The fill uses
`a:tile sx="50000" sy="50000"`. Usage: make_shape_tile.py <out.pptx>. python-pptx and
Pillow are needed only for generation.
"""
from io import BytesIO
import sys

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'


def checkerboard():
    image = Image.new('RGB', (64, 64), (255, 255, 255))
    for y in range(64):
        for x in range(64):
            if (x // 16 + y // 16) % 2 == 0:
                image.putpixel((x, y), (200, 40, 40) if x < 32 else (40, 40, 200))
    out = BytesIO()
    image.save(out, 'PNG')  # Pillow writes no pHYs chunk unless dpi is given
    out.seek(0)
    return out


def main(path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RIGHT_TRIANGLE, Inches(1), Inches(1), Inches(11), Inches(5.5))
    _, rid = slide.part.get_or_add_image_part(checkerboard())
    sp_pr = shape._element.spPr
    for tag in ('solidFill', 'noFill', 'gradFill', 'blipFill', 'pattFill'):
        for old in sp_pr.findall(f'{{{A}}}{tag}'):
            sp_pr.remove(old)
    fill = etree.fromstring(
        f'<a:blipFill xmlns:a="{A}" xmlns:r="{R}"><a:blip r:embed="{rid}"/>'
        '<a:tile tx="0" ty="0" sx="50000" sy="50000" flip="none" algn="tl"/></a:blipFill>')
    sp_pr.insert(list(sp_pr).index(sp_pr.find(f'{{{A}}}prstGeom')) + 1, fill)
    prs.save(path)


if __name__ == '__main__':
    main(sys.argv[1])
