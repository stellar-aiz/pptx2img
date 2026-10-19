"""Independent fixture comparing 0.25-6pt line widths and sm/med/lg arrowheads."""
import sys
from pptx import Presentation
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn
p=Presentation();p.slide_width,p.slide_height=Inches(10),Inches(10)
for size in ['sm','med','lg']:
 s=p.slides.add_slide(p.slide_layouts[6])
 for i,width in enumerate([.25,.5,.75,1,1.5,2,3,4,6]):
  c=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(1),Inches(i+1),Inches(9),Inches(i+1));c.line.width=Pt(width);c.line.color.rgb=RGBColor(0,0,0)
  ln=c._element.spPr.get_or_add_ln();end=OxmlElement('a:tailEnd');end.set('type','triangle');end.set('w',size);end.set('len',size);ln.append(end)
  c._element.remove(c._element.find(qn('p:style')))
p.save(sys.argv[1])
