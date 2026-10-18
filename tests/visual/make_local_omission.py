"""余白中の小さなグラフ省略を再現する自作資料。生成時だけ python-pptx が必要。"""
import json
from pathlib import Path
import sys

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches


def main(out: Path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(10), Inches(5.625)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    data = CategoryChartData()
    data.categories = ["a", "b"]
    data.add_series("series", [1, 2])
    chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(1), Inches(1), Inches(.4), Inches(.4), data).chart
    chart.has_legend = False
    chart.has_title = False
    for axis in (chart.category_axis, chart.value_axis):
        axis.visible = False
    prs.save(out)
    out.with_suffix(".json").write_text(json.dumps({"regions": [
        {"slide": 1, "name": "小グラフ", "rect": [.1, 1 / 5.625, .04, .4 / 5.625]}
    ]}, ensure_ascii=False, indent=2) + "\n")


if __name__ == "__main__":
    main(Path(sys.argv[1]))
