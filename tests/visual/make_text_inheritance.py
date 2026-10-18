"""一般テキストボックスとプレゼン既定書式の自作比較資料。"""
import argparse
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn


def make(output):
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(6), Inches(3.5)
    default = deck._element.find(qn('p:defaultTextStyle'))
    for node in list(default): default.remove(node)
    level = OxmlElement('a:lvl1pPr')
    run = OxmlElement('a:defRPr'); run.set('sz', '1800')
    face = OxmlElement('a:latin'); face.set('typeface', 'Arial'); run.append(face)
    level.append(run); default.append(level)
    # textboxの既定と異なる値。誤ったotherStyleの適用を単独再現する。
    other = deck.slide_masters[0]._element.find(qn('p:txStyles')).find(qn('p:otherStyle'))
    for node in list(other): other.remove(node)
    lev = OxmlElement('a:lvl1pPr'); props = OxmlElement('a:defRPr'); props.set('sz', '2400')
    face = OxmlElement('a:latin'); face.set('typeface', 'Arial'); props.append(face);lev.append(props);other.append(lev)
    for explicit in [False, True]:
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        shape = slide.shapes.add_textbox(Inches(.25), Inches(.2), Inches(3.8), Inches(2.8))
        frame = shape.text_frame; frame.auto_size = MSO_AUTO_SIZE.NONE; frame.word_wrap = True
        frame.margin_top = frame.margin_bottom = frame.margin_left = frame.margin_right = Pt(0)
        for i in range(2):
            p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            p.line_spacing = 1.2; p.space_before = Pt(0); p.space_after = Pt(4)
            r = p.add_run(); r.text = 'Independent sample paragraph with repeated text for line wrapping.'
            if explicit: r.font.size = Pt(18); r.font.name = 'Arial'
    deck.save(output)

if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__); parser.add_argument('output', type=Path)
    make(parser.parse_args().output)
