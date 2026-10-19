"""Independent fixture for color transforms (luminance, tint, shade, saturation, alpha) and theme fillRef."""
from pathlib import Path
import sys
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches
from zipfile import ZipFile,ZIP_DEFLATED
from io import BytesIO
A='http://schemas.openxmlformats.org/drawingml/2006/main'
# Built-in table style ID; the definition below is written for this fixture, mirroring
# the documented region layout (accent header, 20%/40% tinted bands).
STYLE='{5C22544A-7EE6-4342-B048-85BDC9FD1C3A}'

def main(out):
 p=Presentation();p.slide_width,p.slide_height=Inches(10),Inches(5.625)
 slide=p.slides.add_slide(p.slide_layouts[6])
 for row in range(2):
  for i in range(4):
   s=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(.6+i*2.4),Inches(.8+row*2),Inches(1.8),Inches(1.4))
   s.line.fill.background();s._element.remove(s._element.txBody)
   style=s._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}style')
   for e in list(style):style.remove(e)
   fill=OxmlElement('a:fillRef' if row else 'a:solidFill')
   if row:fill.set('idx','1')
   c=OxmlElement('a:schemeClr');c.set('val','accent1');fill.append(c)
   for tag,value in ([('lumMod','50000')],[('lumOff','25000')],[('lumMod','50000'),('lumOff','50000')],[])[i]:
    e=OxmlElement('a:'+tag);e.set('val',value);c.append(e)
   (style if row else s._element.spPr).append(fill)
   if row and i==3:s.fill.solid();s.fill.fore_color.rgb=RGBColor(255,0,0)
 transforms(p)
 buf=BytesIO();p.save(buf)
 bands=''.join(f'<a:{edge}><a:ln w="12700"><a:solidFill><a:schemeClr val="lt1"/></a:solidFill></a:ln></a:{edge}>' for edge in ['left','right','top','bottom','insideH','insideV'])
 def region(name,tint,tx=''):
  t=f'<a:tint val="{tint}"/>' if tint else ''
  return f'<a:{name}>{tx}<a:tcStyle><a:fill><a:solidFill><a:schemeClr val="accent1">{t}</a:schemeClr></a:solidFill></a:fill></a:tcStyle></a:{name}>'
 style=(f'<a:tblStyleLst xmlns:a="{A}" def="{STYLE}"><a:tblStyle styleId="{STYLE}" styleName="Medium Style 2 - Accent 1">'
  f'<a:wholeTbl><a:tcTxStyle><a:schemeClr val="dk1"/></a:tcTxStyle><a:tcStyle><a:tcBdr>{bands}</a:tcBdr><a:fill><a:solidFill><a:schemeClr val="accent1"><a:tint val="20000"/></a:schemeClr></a:solidFill></a:fill></a:tcStyle></a:wholeTbl>'
  +region('band1H','40000')+region('firstRow','','<a:tcTxStyle b="on"><a:schemeClr val="lt1"/></a:tcTxStyle>')+'</a:tblStyle></a:tblStyleLst>')
 with ZipFile(buf) as src,ZipFile(out,'w',ZIP_DEFLATED) as dest:
  for n in src.namelist():dest.writestr(n,style if n=='ppt/tableStyles.xml' else src.read(n))

def transforms(p):
 """Page 2: alpha overlays, tint/shade/satMod swatches, and a table style built from tints."""
 slide=p.slides.add_slide(p.slide_layouts[6])
 def shape(x,y,w,h,color,mods=(),line=None):
  s=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
  s._element.remove(s._element.txBody)
  style=s._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}style')
  s._element.remove(style)
  fill=OxmlElement('a:solidFill' if color else 'a:noFill')
  if color:
   c=OxmlElement('a:schemeClr' if color.startswith('accent') else 'a:srgbClr');c.set('val',color);fill.append(c)
   for tag,value in mods:
    e=OxmlElement('a:'+tag);e.set('val',value);c.append(e)
  s._element.spPr.append(fill)
  ln=OxmlElement('a:ln')
  if line:
   ln.set('w','76200');lf=OxmlElement('a:solidFill');c=OxmlElement('a:srgbClr');c.set('val',line);lf.append(c)
   e=OxmlElement('a:alpha');e.set('val','50000');c.append(e);ln.append(lf)
  else:ln.append(OxmlElement('a:noFill'))
  s._element.spPr.append(ln)
 # Alpha: striped background, then translucent white overlay and translucent line.
 for i,color in enumerate(['C00000','00B050','0070C0','FFC000']):shape(.3+i*.6,.3,.6,2.2,color)
 shape(.5,.6,2.0,1.6,'FFFFFF',[('alpha','76090')])
 shape(.3,1.2,2.4,.5,'',line='000000')
 # Tint, shade, and saturation swatches of accent1.
 for i,mods in enumerate([[],[('tint','20000')],[('tint','40000')],[('tint','75000')],[('shade','50000')],[('shade','75000')],[('satMod','50000')],[('tint','60000'),('satMod','150000')]]):
  shape(3.0+i*.85,.3,.75,.9,'accent1',mods)
 shape(3.0,1.4,6.75,1.1,'accent2',[('alpha','40000')])
 table=slide.shapes.add_table(4,3,Inches(.3),Inches(2.8),Inches(9.4),Inches(2.5)).table
 props=table._tbl.tblPr;props.set('firstRow','1');props.set('bandRow','1')
 for r in range(4):
  for c in range(3):table.cell(r,c).text=''
if __name__=='__main__':main(Path(sys.argv[1]))
