"""Compare independently created table styles with equivalent direct formatting and merged cells."""
from pathlib import Path
from zipfile import ZipFile,ZIP_DEFLATED
from io import BytesIO
import sys
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement
A='http://schemas.openxmlformats.org/drawingml/2006/main'
ID='{EA3C1175-5B30-4D17-AE54-4BB52F704F10}'

def main(out):
 p=Presentation();p.slide_width,p.slide_height=Inches(10),Inches(5.625)
 for explicit in [False,True]:
  slide=p.slides.add_slide(p.slide_layouts[6]);table=slide.shapes.add_table(4,3,Inches(.8),Inches(.7),Inches(8.4),Inches(4.2)).table
  table.cell(0,0).merge(table.cell(0,1))
  props=table._tbl.tblPr;props.set('firstRow','1');props.set('bandRow','1')
  props.find('{'+A+'}tableStyleId').text=ID if not explicit else ''
  if explicit:
   for child in list(props):props.remove(child)
   props.attrib.clear()
  for r in range(4):
   for c in range(3):
    cell=table.cell(r,c)
    if cell.is_spanned:continue
    cell.text=f'項目{r+1}－{c+1}'
    par=cell.text_frame.paragraphs[0];par.font.name='Hiragino Sans';par.font.size=Pt(20)
    for run in par.runs:
     e=OxmlElement('a:ea');e.set('typeface','Hiragino Sans');run._r.get_or_add_rPr().append(e)
    color='245A51' if r==0 else ('D2EDFE' if r%2 else 'FFFFFF')
    if explicit or (r,c)==(3,2):
     cell.fill.solid();cell.fill.fore_color.rgb=RGBColor.from_string('C37B49' if (r,c)==(3,2) else color)
     par.font.bold=r==0;par.font.color.rgb=RGBColor.from_string('FFFFFF' if r==0 else '000000')
     for edge in ['lnL','lnR','lnT','lnB']:
      line=OxmlElement('a:'+edge);line.set('w','12700');f=OxmlElement('a:solidFill');rgb=OxmlElement('a:srgbClr');rgb.set('val','888888');f.append(rgb);line.append(f);cell._tc.get_or_add_tcPr().append(line)
 buf=BytesIO();p.save(buf)
 borders=''.join(f'<a:{edge}><a:ln w="12700"><a:solidFill><a:srgbClr val="888888"/></a:solidFill></a:ln></a:{edge}>' for edge in ['left','right','top','bottom','insideH','insideV'])
 def region(name,color,tx=''):return f'<a:{name}>{tx}<a:tcStyle><a:fill><a:solidFill><a:srgbClr val="{color}"/></a:solidFill></a:fill></a:tcStyle></a:{name}>'
 style=f'<a:tblStyleLst xmlns:a="{A}"><a:tblStyle styleId="{ID}" styleName="自作比較"><a:wholeTbl><a:tcTxStyle b="off"><a:srgbClr val="000000"/></a:tcTxStyle><a:tcStyle><a:tcBdr>{borders}</a:tcBdr><a:fill><a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill></a:fill></a:tcStyle></a:wholeTbl>'+region('firstRow','245A51','<a:tcTxStyle b="on"><a:srgbClr val="FFFFFF"/></a:tcTxStyle>')+region('band1H','D2EDFE')+'</a:tblStyle></a:tblStyleLst>'
 with ZipFile(buf) as src,ZipFile(out,'w',ZIP_DEFLATED) as dest:
  for n in src.namelist():dest.writestr(n,style if n=='ppt/tableStyles.xml' else src.read(n))
if __name__=='__main__':main(Path(sys.argv[1]))
