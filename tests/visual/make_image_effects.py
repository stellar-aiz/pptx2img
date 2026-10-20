"""Independent fixture: picture effects, tiled picture fills, and fill rectangles.

Every slide shows the same generated picture (a hue sweep over a grey ramp with a white
border, drawn here with Pillow) twice: unchanged on the left, and with one `a:blip`
effect or picture fill on the right, over a checkered backdrop so transparency shows.

1. `a:alphaModFix amt="50000"`        7. `a:duotone` (dark blue to light yellow)
2. `a:grayscl`                        8. `a:clrChange` white -> transparent
3. `a:biLevel thresh="50000"`         9. `a:tile` fills: scale, alignment, flips
4. `a:lum bright/contrast`           10. `a:stretch/a:fillRect` insets
5. `a14:imgProps` saturation 0       11. metadata only (`a14:useLocalDpi`) and
6. `a14:imgProps` brightnessContrast     `a14:imgProps` backgroundRemoval (named, not drawn)

Writes `<out>` and per-picture regions `<out>.json`. python-pptx and Pillow are needed
only for generation.
"""
from io import BytesIO
import colorsys
import json
from pathlib import Path
import sys

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
A14 = 'http://schemas.microsoft.com/office/drawing/2010/main'
IMG_PROPS_EXT = '{BEBA8EAE-BF5A-486C-A8C5-ECC9F3942E4B}'
LOCAL_DPI_EXT = '{28A0092B-C50C-407E-A947-70E740481C1C}'

SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)
PIC_W, PIC_H = Inches(4.8), Inches(3.6)
LEFT_X, RIGHT_X, TOP_Y = Inches(1.0), Inches(7.4), Inches(2.4)


def picture_png():
    """160x120 px: hues across, a dark-to-light ramp down, a white border, 96 dpi."""
    width, height = 160, 120
    image = Image.new('RGB', (width, height), (255, 255, 255))
    for x in range(8, width - 8):
        for y in range(8, height - 8):
            hue = (x - 8) / (width - 16)
            value = 0.25 + 0.75 * (y - 8) / (height - 17)
            r, g, b = colorsys.hsv_to_rgb(hue, 0.85, value)
            image.putpixel((x, y), (round(r * 255), round(g * 255), round(b * 255)))
    data = BytesIO()
    image.save(data, 'PNG', dpi=(96, 96))
    data.seek(0)
    return data


def checker(slide, x, y, w, h):
    """An 8x6 checkerboard of blue squares behind the right picture."""
    cols, rows = 8, 6
    for i in range(cols):
        for j in range(rows):
            if (i + j) % 2:
                continue
            shape = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, x + w * i // cols, y + h * j // rows, w // cols, h // rows
            )
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor(0x30, 0x60, 0x90)
            shape.line.fill.background()


def label(slide, text):
    box = slide.shapes.add_textbox(Inches(1.0), Inches(0.6), Inches(11.3), Inches(1.0))
    box.text_frame.text = text
    box.text_frame.paragraphs[0].runs[0].font.size = Pt(28)


def sub(parent, tag, ns=A, **attrs):
    return etree.SubElement(parent, f'{{{ns}}}{tag}', {k: str(v) for k, v in attrs.items()})


def srgb(parent, value, alpha=None):
    color = sub(parent, 'srgbClr', val=value)
    if alpha is not None:
        sub(color, 'alpha', val=alpha)
    return color


def extension(blip, uri):
    ext_list = blip.find(f'{{{A}}}extLst')
    if ext_list is None:
        ext_list = sub(blip, 'extLst')
    return sub(ext_list, 'ext', uri=uri)


def office_effect(blip, tag, **attrs):
    ext = extension(blip, IMG_PROPS_EXT)
    props = etree.SubElement(ext, f'{{{A14}}}imgProps', nsmap={'a14': A14})
    layer = sub(props, 'imgLayer', A14)
    effect = sub(layer, 'imgEffect', A14)
    sub(effect, tag, A14, **attrs)


def add_pair(slide, title, edit_blip=None, edit_fill=None):
    label(slide, title)
    slide.shapes.add_picture(picture_png(), LEFT_X, TOP_Y, PIC_W, PIC_H)
    checker(slide, RIGHT_X, TOP_Y, PIC_W, PIC_H)
    pic = slide.shapes.add_picture(picture_png(), RIGHT_X, TOP_Y, PIC_W, PIC_H)
    blip_fill = pic._element.blipFill
    if edit_blip:
        edit_blip(blip_fill.blip)
    if edit_fill:
        for old in blip_fill.findall(f'{{{A}}}stretch'):
            blip_fill.remove(old)
        edit_fill(blip_fill)


def tile(**attrs):
    return lambda fill: sub(fill, 'tile', **attrs)


def main(out):
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    pages = [
        ('1. alphaModFix 50%', lambda b: sub(b, 'alphaModFix', amt=50000), None),
        ('2. grayscl', lambda b: sub(b, 'grayscl'), None),
        ('3. biLevel 50%', lambda b: sub(b, 'biLevel', thresh=50000), None),
        ('4. lum bright +20% contrast +30%', lambda b: sub(b, 'lum', bright=20000, contrast=30000), None),
        ('5. a14 saturation 0', lambda b: office_effect(b, 'saturation', sat=0), None),
        ('6. a14 brightnessContrast -20% / +40%',
         lambda b: office_effect(b, 'brightnessContrast', bright=-20000, contrast=40000), None),
        ('7. duotone (approximation)', lambda b: (
            lambda d: (srgb(d, '1F3864'), srgb(d, 'FFF2CC')))(sub(b, 'duotone')), None),
        ('8. clrChange white to transparent (approximation)', lambda b: (
            lambda c: (srgb(sub(c, 'clrFrom'), 'FFFFFF'), srgb(sub(c, 'clrTo'), 'FFFFFF', 0)))(
            sub(b, 'clrChange')), None),
        ('9a. tile 50%, top-left', None, tile(tx=0, ty=0, sx=50000, sy=50000, flip='none', algn='tl')),
        ('9b. tile 40%, centred, offset', None,
         tile(tx=Inches(0.2), ty=Inches(0.1), sx=40000, sy=40000, flip='none', algn='ctr')),
        ('9c. tile 50%, flip xy, bottom-right', None,
         tile(tx=0, ty=0, sx=50000, sy=50000, flip='xy', algn='br')),
        ('10. fillRect inset 10% / 20%', None,
         lambda f: sub(sub(f, 'stretch'), 'fillRect', l=10000, t=20000, r=10000, b=20000)),
        ('11. useLocalDpi + backgroundRemoval (drawn unchanged)', lambda b: (
            etree.SubElement(extension(b, LOCAL_DPI_EXT), f'{{{A14}}}useLocalDpi', {'val': '0'},
                             nsmap={'a14': A14}),
            office_effect(b, 'backgroundRemoval', t=0, b=100000, l=0, r=100000)), None),
    ]
    regions = []
    for index, (title, edit_blip, edit_fill) in enumerate(pages, start=1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        add_pair(slide, title, edit_blip, edit_fill)
        for name, x in (('original', LEFT_X), ('effect', RIGHT_X)):
            regions.append({
                'slide': index,
                'name': name,
                'rect': [
                    round(int(x) / SLIDE_W, 4),
                    round(int(TOP_Y) / SLIDE_H, 4),
                    round(int(PIC_W) / SLIDE_W, 4),
                    round(int(PIC_H) / SLIDE_H, 4),
                ],
            })
    prs.save(out)
    Path(out).with_suffix('.json').write_text(json.dumps({'regions': regions}, indent=2) + '\n')


if __name__ == '__main__':
    main(Path(sys.argv[1]))
