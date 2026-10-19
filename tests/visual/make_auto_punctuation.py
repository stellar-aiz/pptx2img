"""Automatic punctuation compression fixtures at different widths; python-pptx is needed only for regeneration.

Includes 100/110/120/130/140/150pt, preserving boundary line-break differences such as those at 120pt.
All runs leave spc unspecified. Do not confuse automatic compression with explicit spacing.
"""
import argparse
from pathlib import Path
from pptx import Presentation
from pptx.util import Pt, Inches
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.oxml.xmlchemy import OxmlElement


def make(output: Path):
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(8), Inches(6)

    def box(slide, text, x, y, width, hanging=False):
        shape = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(width), Pt(120))
        frame = shape.text_frame
        frame.auto_size, frame.word_wrap = MSO_AUTO_SIZE.NONE, True
        frame.margin_left = frame.margin_right = 0
        frame.margin_top = frame.margin_bottom = 0
        para = frame.paragraphs[0]
        para._p.get_or_add_pPr().set('hangingPunct', '1' if hanging else '0')
        run = para.add_run()
        run.text = text
        run.font.name, run.font.size = 'Hiragino Sans', Pt(24)
        props = run._r.get_or_add_rPr()
        props.set('lang', 'ja-JP')
        east_asian = OxmlElement('a:ea')
        east_asian.set('typeface', 'Hiragino Sans')
        props.append(east_asian)

    def width_cases(text, hanging=False):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        for i, width in enumerate((100, 110, 120, 130, 140, 150)):
            box(slide, text, 30 + i % 2 * 280, 10 + i // 2 * 130, width, hanging)

    for text in ('あ、いう、え', 'あ「い」う、', '「「あ」」い',
                 'あいうえおか', 'あ・い：う；', 'あ。いう。え'):
        width_cases(text)
    # Include dense punctuation so an implementation ignoring the setting cannot pass on whole-page averages.
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    for i in range(6):
        box(slide, 'あ、い、う、え、お、か、き、', 30, 10 + i * 65, 310)
    # Configure compression and hanging punctuation separately.
    width_cases('あ、いう、え', hanging=True)
    deck.save(output)


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('output', type=Path)
    make(parser.parse_args().output)
