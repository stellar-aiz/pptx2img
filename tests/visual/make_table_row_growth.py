"""Create a self-made deck whose table rows are stored shorter than their wrapped text.

Slide 1: a 3x2 table with 0.4in rows whose middle row holds a long paragraph that wraps
to several lines; Office and LibreOffice grow that row and push the last row down.
Slide 2: rows that fit keep their stored heights although the graphicFrame height is
twice their sum (the frame height does not stretch rows).
"""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

COLORS = ["FFD9D9", "D9FFD9", "D9D9FF"]
LONG = ("This cell holds a long sentence that wraps over several lines, so its row "
        "must grow taller than the stored height to keep the text inside the cell.")


def table(slide, texts, row_height, frame_height):
    frame = slide.shapes.add_table(3, 2, Inches(1), Inches(1), Inches(8), Inches(frame_height))
    t = frame.table
    for r in range(3):
        t.rows[r].height = Inches(row_height)
        for c in range(2):
            cell = t.cell(r, c)
            cell.text = texts[r][c]
            font = cell.text_frame.paragraphs[0].font
            font.name, font.size = "Arial", Pt(18)
            font.color.rgb = RGBColor(0, 0, 0)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(COLORS[r])
    # python-pptx resizes the frame to the row sum; restore the requested frame height.
    frame.height = Inches(frame_height)


def main(out):
    p = Presentation()
    p.slide_width, p.slide_height = Inches(10), Inches(5.625)
    grow = p.slides.add_slide(p.slide_layouts[6])
    table(grow, [["Row 1", "Short"], ["Row 2", LONG], ["Row 3", "Short"]], 0.4, 1.2)
    frame = p.slides.add_slide(p.slide_layouts[6])
    table(frame, [["Row 1", "Fits"], ["Row 2", "Fits"], ["Row 3", "Fits"]], 0.5, 3.0)
    p.save(out)


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "examples/table-row-growth.pptx")
