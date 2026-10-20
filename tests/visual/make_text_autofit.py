"""Independent fixture isolating saved text scaling and shape autofit."""
import argparse
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn


def make(output):
    deck = Presentation(); deck.slide_width, deck.slide_height = Inches(6), Inches(3.5)
    for scale in [100000, 55000, 92500]:
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        shape = slide.shapes.add_textbox(Inches(.3), Inches(.4), Inches(5.4), Inches(2.5))
        frame = shape.text_frame; frame.auto_size=MSO_AUTO_SIZE.NONE;frame.word_wrap=True
        frame.margin_top=frame.margin_bottom=frame.margin_left=frame.margin_right=Pt(0)
        body=frame._txBody.bodyPr
        for node in list(body): body.remove(node)
        auto=OxmlElement('a:normAutofit');auto.set('fontScale',str(scale));auto.set('lnSpcReduction','20000');body.append(auto)
        for i in range(3):
            p=frame.paragraphs[0] if i==0 else frame.add_paragraph();p.line_spacing=1.2
            p.space_before=Pt(0);p.space_after=Pt(0)
            run=p.add_run();run.text='Auto fit sample ABC 123';run.font.size=Pt(24);run.font.name='Arial';run.font.color.rgb=RGBColor(0,0,0)
    slide=deck.slides.add_slide(deck.slide_layouts[6])
    for i,anchor in enumerate(['t','ctr','b']):
        shape=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(.3+i*1.9), Inches(1.2), Inches(1.6), Inches(.4))
        preset = shape._element.find(qn("p:style"))
        if preset is not None: shape._element.remove(preset)
        shape.line.fill.background()
        shape.fill.solid();shape.fill.fore_color.rgb=RGBColor(230,240,255)
        frame=shape.text_frame;frame.auto_size=MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT;frame.word_wrap=True
        frame.margin_top=frame.margin_bottom=Pt(4);frame.margin_left=frame.margin_right=Pt(4)
        frame._txBody.bodyPr.set('anchor',anchor)
        for j in range(3):
            p=frame.paragraphs[0] if j==0 else frame.add_paragraph();p.line_spacing=1.0;p.alignment=PP_ALIGN.LEFT
            run=p.add_run();run.text='Expand '+str(j+1);run.font.size=Pt(18);run.font.name='Arial';run.font.color.rgb=RGBColor(0,0,0)
    # Unwrapped shape autofit: saved widths that are too narrow (top row) and too wide
    # (bottom row) for left, centred, and right-aligned text.
    slide=deck.slides.add_slide(deck.slide_layouts[6])
    for row,width in enumerate([.6,1.7]):
        for i,align in enumerate([PP_ALIGN.LEFT,PP_ALIGN.CENTER,PP_ALIGN.RIGHT]):
            shape=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(.3+i*1.9+(1.7-width)/2),Inches(.6+row*1.4),Inches(width),Inches(.5))
            preset=shape._element.find(qn("p:style"))
            if preset is not None: shape._element.remove(preset)
            shape.line.color.rgb=RGBColor(40,80,160);shape.line.width=Pt(1)
            shape.fill.solid();shape.fill.fore_color.rgb=RGBColor(230,240,255)
            frame=shape.text_frame;frame.auto_size=MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT;frame.word_wrap=False
            frame.margin_top=frame.margin_bottom=Pt(4);frame.margin_left=frame.margin_right=Pt(6)
            p=frame.paragraphs[0];p.alignment=align
            run=p.add_run();run.text='Label '+'LCR'[i];run.font.size=Pt(18);run.font.name='Arial';run.font.color.rgb=RGBColor(0,0,0)
    deck.save(output)

if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__);parser.add_argument('output',type=Path)
    make(parser.parse_args().output)
