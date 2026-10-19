"""Independent fixture: pictures whose references cannot be resolved the usual way.

One slide, four generated PNG pictures, each with a caption:

1. a control picture with an ordinary image relationship;
2. a picture whose relationship type is rewritten to `oleObject` (as some third-party
   writers do); pptx2img loads the part because the picture references it and renders it;
3. a picture whose `r:embed` points at a relationship id that does not exist;
4. a picture whose crop (`srcRect`) leaves no visible area.

Pictures 3 and 4 are omitted with `omitted` diagnostics; the rest of the slide renders.
python-pptx and Pillow are needed only for generation.
"""
from io import BytesIO
from pathlib import Path
import sys

from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt

OLE_OBJECT = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/oleObject'


def png(color):
    data = BytesIO()
    Image.new('RGB', (160, 120), color).save(data, 'PNG')
    data.seek(0)
    return data


def main(out):
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(10), Inches(5.625)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    cases = [
        ('Control', (0, 102, 204)),
        ('oleObject type', (0, 153, 51)),
        ('Dangling r:embed', (204, 51, 0)),
        ('Empty crop', (153, 0, 153)),
    ]
    for index, (caption, color) in enumerate(cases):
        left = Inches(0.4 + index * 2.4)
        pic = slide.shapes.add_picture(png(color), left, Inches(1.2), Inches(2), Inches(1.5))
        box = slide.shapes.add_textbox(left, Inches(3), Inches(2.2), Inches(0.6))
        box.text_frame.text = caption
        box.text_frame.paragraphs[0].runs[0].font.size = Pt(16)
        rid = pic._element.blip_rId
        if index == 1:
            slide.part.rels[rid]._reltype = OLE_OBJECT
        elif index == 2:
            pic._element.blipFill.blip.rEmbed = 'rId999'
        elif index == 3:
            pic.crop_left, pic.crop_right = 0.6, 0.5
    deck.save(out)


if __name__ == '__main__':
    main(Path(sys.argv[1]))
