"""保存経路と接続点参照の有無を比較する自作資料。"""
import sys
from make_lines import line
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches
p=Presentation();p.slide_width,p.slide_height=Inches(5),Inches(4)
for connect in [False,True]:
 s=p.slides.add_slide(p.slide_layouts[6])
 a=s.shapes.add_shape(MSO_SHAPE.PENTAGON,Inches(1.45),Inches(.25),Inches(1),Inches(.5))
 b=s.shapes.add_shape(MSO_SHAPE.PENTAGON,Inches(2.45),Inches(2.5),Inches(1),Inches(.5))
 c=line(s,1.2,1.5,2,.5,'bentConnector2',tail='triangle')
 c.rotation=270;c._element.spPr.xfrm.set('flipH','1')
 if connect:
  for tag,id,index in [('stCxn',a.shape_id,2),('endCxn',b.shape_id,1)]:
   node=OxmlElement('a:'+tag);node.set('id',str(id));node.set('idx',str(index));c._element.nvCxnSpPr.cNvCxnSpPr.append(node)
p.save(sys.argv[1])
