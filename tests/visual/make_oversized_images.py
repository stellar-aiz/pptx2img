"""Independent fixture: pictures above the embed limit (4400x4400 PNG, 5800x3456 JPEG).

rhodopsin downscales both with `approximation` diagnostics instead of aborting. The
script also writes a same-layout deck without pictures, for peak-memory comparison.
python-pptx and Pillow are needed only for generation.
"""
from io import BytesIO
from pathlib import Path
import sys

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches, Pt


def picture(size, fmt):
    width, height = size
    image = Image.new('RGB', size, '#245A51')
    draw = ImageDraw.Draw(image)
    draw.rectangle((width // 10, height // 10, width // 2, height * 9 // 10), fill='#C37B49')
    draw.ellipse((width * 6 // 10, height // 4, width * 9 // 10, height * 3 // 4), fill='#EEEEEE')
    data = BytesIO()
    image.save(data, format=fmt, **({'quality': 90} if fmt == 'JPEG' else {}))
    return data


def deck(out, with_pictures):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(10), Inches(5.625)
    for label, size, fmt in [('PNG 4400x4400', (4400, 4400), 'PNG'),
                             ('JPEG 5800x3456', (5800, 3456), 'JPEG')]:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(.5), Inches(.3), Inches(9), Inches(.8))
        box.text_frame.text = label
        box.text_frame.paragraphs[0].runs[0].font.size = Pt(24)
        if with_pictures:
            height = Inches(4.2)
            width = int(height * size[0] / size[1])
            slide.shapes.add_picture(picture(size, fmt), Inches(.5), Inches(1.2), width, height)
    prs.save(out)


if __name__ == '__main__':
    out = Path(sys.argv[1])
    deck(out, True)
    deck(out.with_name(out.stem + '-baseline.pptx'), False)
