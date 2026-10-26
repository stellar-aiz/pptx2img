"""Compare production PDF output with PNG and record extractor differences.

Usage: .venv-visual/bin/python tests/visual/pdf_output.py BINARY FONT_FILE FONT_FAMILY OUT
The output directory must not exist. No font files are copied into the results.
"""
import json
from pathlib import Path
import subprocess
import sys
import struct
from io import BytesIO
from PIL import Image, ImageDraw
import pypdfium2 as pdfium
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from compare import compare_images, passes


def run(binary, font, family, out):
    out.mkdir()
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(10), Inches(5.625)
    slide = deck.slides.add_slide(deck.slide_layouts[6])

    def label(slide, text, x, y, width, size):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(.6))
        paragraph = box.text_frame.paragraphs[0]
        paragraph.text = text
        paragraph.font.name, paragraph.font.size = family, Pt(size)
        paragraph.font.color.rgb = RGBColor(20, 30, 40)

    expected = ['日本語の計画書', '工期は12か月、予算は3,500万円。', '項目', '数量', '橋梁', '12', '舗装', '3500', 'Second page 123']
    label(slide, expected[0], .5, .35, 8, 24)
    label(slide, expected[1], .5, 1, 8, 12)
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(.5), Inches(1.8), Inches(3), Inches(.7))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(30, 100, 170)
    shape.line.color.rgb = RGBColor(10, 40, 80)
    table = slide.shapes.add_table(3, 2, Inches(.5), Inches(2.9), Inches(5), Inches(1.5)).table
    for row, values in enumerate([['項目', '数量'], ['橋梁', '12'], ['舗装', '3500']]):
        for col, value in enumerate(values):
            cell = table.cell(row, col)
            cell.text = value
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.name, paragraph.font.size = family, Pt(14)
    image = Image.new('RGB', (120, 80), 'white')
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, 59, 39), fill='red')
    draw.rectangle((60, 40, 119, 79), fill='blue')
    data = BytesIO()
    image.save(data, format='PNG')
    data.seek(0)
    slide.shapes.add_picture(data, Inches(6.5), Inches(2), width=Inches(2.4))
    jpeg = BytesIO()
    image.save(jpeg, format='JPEG')
    jpeg.seek(0)
    slide.shapes.add_picture(jpeg, Inches(6.5), Inches(3.7), width=Inches(1.2))
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    label(slide, expected[-1], .5, 1, 8, 18)
    hidden = deck.slides.add_slide(deck.slide_layouts[6])
    hidden._element.set('show', '0')
    label(hidden, 'HIDDEN SENTINEL', .5, 1, 8, 18)
    fixture = out / 'supported.pptx'
    deck.save(fixture)
    png = out / 'png'
    pdf = out / 'output.pdf'
    common = [str(binary), str(fixture), '--no-system-fonts', '--font-file', str(font), '--font-fallback']
    subprocess.run([*common, '--out', str(png)], check=True, capture_output=True)
    result = subprocess.run([*common, '--format', 'pdf', '--out', str(pdf)], check=True, capture_output=True, text=True)
    (out / 'result.yaml').write_text(result.stdout)
    extraction = {}
    for mode, flags in [('default', []), ('layout', ['-layout']), ('raw', ['-raw'])]:
        text = subprocess.check_output(['pdftotext', *flags, pdf, '-'], text=True)
        extraction[mode] = {'text': text, 'exact_strings': {value: value in text for value in expected}}
        assert 'HIDDEN' not in text
    document = pdfium.PdfDocument(pdf)
    assert len(document) == 2
    metrics = []
    for index, page in enumerate(document):
        assert tuple(page.get_size()) == (720.0, 405.0)
        reference = Image.open(png / f'slide-{index+1:03}.png').convert('RGB')
        actual = page.render(scale=1600 / 720).to_pil().convert('RGB')
        actual.save(out / f'pdfium-{index+1}.png')
        measured, difference = compare_images(actual, reference, 16)
        difference.save(out / f"difference-{index+1}.png")
        measured['passed'] = passes(measured, .02, .05)
        measured['pdfium_text'] = page.get_textpage().get_text_range()
        measured['regions'] = []
        if index == 0:
            for name, box in [('title', (60, 60, 1000, 150)), ('sentence', (60, 150, 1100, 220)), ('table', (60, 450, 1000, 720)), ('image', (1030, 310, 1450, 600))]:
                region, _ = compare_images(actual.crop(box), reference.crop(box), 16)
                region.update(name=name, passed=passes(region, .02, .05))
                measured['regions'].append(region)
        metrics.append(measured)
    combined = '\n'.join(m['pdfium_text'] for m in metrics)
    extraction['pdfium'] = {'text': combined, 'exact_strings': {value: value in combined for value in expected}}
    summary = {'extraction': extraction, 'pages': 2, 'hidden_slide_skipped': True,
               'pdf_bytes': pdf.stat().st_size, 'metrics': metrics}
    (out / 'summary.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2))
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    assert all(m['passed'] and all(r['passed'] for r in m['regions']) for m in metrics), 'Visual differences exceed existing thresholds'
    assert all(extraction['pdfium']['exact_strings'].values()), 'PDFium must recover fixture text; other extractor differences are recorded'


if __name__ == '__main__':
    run(Path(sys.argv[1]).resolve(), Path(sys.argv[2]).resolve(), sys.argv[3], Path(sys.argv[4]).resolve())
