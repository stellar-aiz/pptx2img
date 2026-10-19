"""Create a self-made deck whose table has a row stored with h="0" (a minimum height of zero)."""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor


def main(out):
    p = Presentation()
    p.slide_width, p.slide_height = Inches(10), Inches(5.625)
    slide = p.slides.add_slide(p.slide_layouts[6])
    table = slide.shapes.add_table(3, 2, Inches(1), Inches(1), Inches(8), Inches(3)).table
    for r in range(3):
        for c in range(2):
            cell = table.cell(r, c)
            cell.text = f"Row {r + 1} Col {c + 1}"
            cell.text_frame.paragraphs[0].font.size = Pt(20)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(["FFD9D9", "D9FFD9", "D9D9FF"][r])
    table.rows[0].height = Inches(1)
    table.rows[1].height = 0
    table.rows[2].height = Inches(1)
    p.save(out)


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "examples/table-zero-row-height.pptx")
