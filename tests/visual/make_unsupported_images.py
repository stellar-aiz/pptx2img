"""Independent fixture: EMF pictures on a slide and in a shared slide layout.

The EMF is a minimal metafile (one red rectangle) written byte by byte here. pptx2img
does not render EMF; it should omit each picture with a diagnostic and still render
every slide. python-pptx is needed only for generation.
"""
from io import BytesIO
from pathlib import Path
import struct
import sys

from pptx import Presentation
from pptx.opc.packuri import PackURI
from pptx.oxml.shapes.picture import CT_Picture
from pptx.util import Inches, Pt


def emf(color=0x000000FF):
    """EMR_HEADER, a solid brush, a 100x100 rectangle, and EMR_EOF (COLORREF is 0x00BBGGRR)."""
    records = [
        struct.pack('<II', 39, 24) + struct.pack('<IIII', 1, 0, color, 0),  # CREATEBRUSHINDIRECT
        struct.pack('<II', 37, 12) + struct.pack('<I', 1),  # SELECTOBJECT
        struct.pack('<II', 43, 24) + struct.pack('<iiii', 0, 0, 100, 100),  # RECTANGLE
        struct.pack('<II', 14, 20) + struct.pack('<III', 0, 16, 20),  # EOF
    ]
    body = b''.join(records)
    total = 88 + len(body)
    header = struct.pack(
        '<II4i4iIIIIHHIIIiiii',
        1, 88,
        0, 0, 100, 100,  # bounds (device units)
        0, 0, 2646, 2646,  # frame (0.01 mm)
        0x464D4520, 0x10000, total, len(records) + 1,
        2, 0, 0, 0, 0,
        1920, 1080, 508, 286,
    )
    return header + body


def as_emf(part, name):
    # Pillow reports EMF as "WMF"; keep the part name and content type faithful to the bytes.
    part.partname = PackURI(f'/ppt/media/{name}.emf')
    part._content_type = 'image/x-emf'


def main(out):
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(10), Inches(5.625)
    layout = deck.slide_layouts[6]
    # A layout logo appears on every slide that uses the layout.
    part, rid = layout.part.get_or_add_image_part(BytesIO(emf(0x00FF0000)))
    as_emf(part, 'layout-logo')
    logo = CT_Picture.new_pic(900, 'EMF logo', '', rid,
                              Inches(8.8), Inches(.2), Inches(1), Inches(1))
    layout.shapes._spTree.append(logo)
    for index in range(2):
        slide = deck.slides.add_slide(layout)
        box = slide.shapes.add_textbox(Inches(.5), Inches(.5), Inches(6), Inches(1))
        box.text_frame.word_wrap = True
        box.text_frame.text = f'Slide {index + 1}'
        box.text_frame.paragraphs[0].runs[0].font.size = Pt(24)
        if index == 0:
            pic = slide.shapes.add_picture(BytesIO(emf()), Inches(1), Inches(2),
                                           Inches(2), Inches(2))
            as_emf(slide.part.related_part(pic._element.blip_rId), 'slide-picture')
    deck.save(out)


if __name__ == '__main__':
    main(Path(sys.argv[1]))
