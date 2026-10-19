"""Independent fixture for straight lines, bent lines, dashes, and arrowheads."""
from pathlib import Path
import sys
from pptx import Presentation
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

def line(slide, x,y,w,h,kind='straightConnector1',dash='solid',head='none',tail='none',size='med'):
    s=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x),Inches(y),Inches(x+w),Inches(y+h))
    s._element.spPr.prstGeom.set('prst',kind)
    s.line.color.rgb=RGBColor.from_string('245A51');s.line.width=Pt(3)
    ln=s._element.spPr.get_or_add_ln();ln.set('cap','flat')
    for tag,attrs in [('prstDash',{'val':dash}),('headEnd',{'type':head,'w':size,'len':size}),('tailEnd',{'type':tail,'w':size,'len':size})]:
        e=OxmlElement('a:'+tag)
        for k,v in attrs.items():e.set(k,v)
        ln.append(e)
    style=s._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}style')
    if style is not None:s._element.remove(style)
    return s

def main(out):
    p=Presentation();p.slide_width,p.slide_height=Inches(10),Inches(5.625)
    slide=p.slides.add_slide(p.slide_layouts[6])
    for i in range(4):
        s=line(slide,.8+i*2.3,1,1.5,2,'bentConnector2' if i%2 else 'straightConnector1',head='triangle',tail='arrow')
        if i>1:s._element.spPr.xfrm.set('flipH','1');s.rotation=20
    slide=p.slides.add_slide(p.slide_layouts[6])
    for i,dash in enumerate(['solid','dot','dash','lgDash','dashDot','lgDashDot','lgDashDotDot','sysDot','sysDash','sysDashDot','sysDashDotDot']):
        line(slide,.6,.4+i*.45,8.5,0,dash=dash)
    slide=p.slides.add_slide(p.slide_layouts[6])
    for i,kind in enumerate(['triangle','stealth','arrow','diamond','oval']):
        for j,size in enumerate(['sm','med','lg']):line(slide,.7+j*3.1,.7+i*.95,2,0,head=kind,tail=kind,size=size)
    p.save(out)
if __name__=='__main__':main(Path(sys.argv[1]))
