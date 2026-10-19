"""Independent fixture for reference font substitution: Japanese text in families that are
usually missing from the comparison machine. python-pptx is needed only for generation."""
from pathlib import Path
import sys

from pptx import Presentation
from pptx.util import Inches, Pt

FAMILIES = ["Meiryo UI", "游ゴシック", "ＭＳ Ｐゴシック", "ＭＳ 明朝", "Calibri", "Yu Mincho"]
TEXT = "日本語の本文サンプル Sample text 0123"


def main(out: Path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(10), Inches(5.625)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for index, family in enumerate(FAMILIES):
        box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3 + index * 0.85), Inches(9), Inches(0.7))
        run = box.text_frame.paragraphs[0].add_run()
        run.text = f"{TEXT} ({family})"
        run.font.name = family
        run.font.size = Pt(22)
        # Set the East Asian typeface too, so CJK glyphs resolve through the same family.
        rpr = run._r.get_or_add_rPr()
        ea = rpr.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}ea", {"typeface": family})
        rpr.append(ea)
    prs.save(out)


if __name__ == "__main__":
    main(Path(sys.argv[1]))
