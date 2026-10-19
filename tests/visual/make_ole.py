"""Compare OLE objects with independently created PNG/JPEG previews against regular images."""
from io import BytesIO
from pathlib import Path
import sys
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches


def main(out):
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(10), Inches(5.625)
    image = Image.new('RGB', (300, 180), '#245A51')
    draw = ImageDraw.Draw(image)
    draw.rectangle((20, 20, 130, 160), fill='#C37B49')
    draw.ellipse((160, 40, 270, 150), fill='#EEEEEE')
    for fmt in ['PNG', 'JPEG']:
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        data = BytesIO()
        image.save(data, format=fmt)
        raw = data.getvalue()
        # OLE payload is independently generated invalid data; preview tests require no execution or parsing.
        slide.shapes.add_ole_object(
            BytesIO(b'own inert OLE fixture'), 'Own.Document',
            Inches(.7), Inches(1.4), Inches(3.6), Inches(2.16),
            icon_file=BytesIO(raw), icon_width=Inches(3.6), icon_height=Inches(2.16),
        )
        slide.shapes.add_picture(BytesIO(raw), Inches(5.5), Inches(1.4), Inches(3.6), Inches(2.16))
    deck.save(out)


if __name__ == '__main__':
    main(Path(sys.argv[1]))
