"""Independent shapes for direction, blur, transforms, and shadow parameters matching source page 10."""
from pathlib import Path
import sys
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches,Pt

def main(out):
 p=Presentation();p.slide_width,p.slide_height=Inches(10),Inches(5.625)
 for page in range(2):
  slide=p.slides.add_slide(p.slide_layouts[6])
  for i in range(4):
   s=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(.8+(i%2)*4.5),Inches(.8+(i//2)*2.4),Inches(3),Inches(1.3))
   if page:s.fill.background()
   else:s.fill.solid();s.fill.fore_color.rgb=RGBColor.from_string('245A51')
   s.line.color.rgb=RGBColor.from_string('C37B49');s.line.width=Pt(3)
   s._element.remove(s._element.txBody)
   st=s._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}style');s._element.remove(st)
   if i>1:s.rotation=15;s._element.spPr.xfrm.set('flipH','1')
   lst=OxmlElement('a:effectLst');sh=OxmlElement('a:outerShdw')
   for k,v in {'blurRad':str(50800 if page else i*35000),'dist':str(38100 if page else 100000),'dir':str(2700000 if page else i*5400000),'algn':'tl','rotWithShape':'0' if i%2==0 else '1'}.items():sh.set(k,v)
   c=OxmlElement('a:prstClr');c.set('val','black');a=OxmlElement('a:alpha');a.set('val','40000');c.append(a);sh.append(c);lst.append(sh);s._element.spPr.append(lst)
 p.save(out)
if __name__=='__main__':main(Path(sys.argv[1]))
