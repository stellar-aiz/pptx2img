"""Independent fixture: bulleted list with empty and whitespace-only paragraphs."""
import copy
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement

p = Presentation()
p.slide_width, p.slide_height = Inches(10), Inches(5.625)
s = p.slides.add_slide(p.slide_layouts[6])
tf = s.shapes.add_textbox(Inches(1), Inches(0.6), Inches(8), Inches(4)).text_frame
tf.word_wrap = True


def bullet(par):
    ppr = par._p.get_or_add_pPr()
    ppr.set("marL", "342900")
    ppr.set("indent", "-342900")
    ppr.append(OxmlElement("a:buChar"))
    ppr[-1].set("char", "•")


# kind: "text", "empty" (no runs), "space" (whitespace-only run)
for i, (kind, text) in enumerate([
    ("text", "First item"), ("text", "Second item"), ("empty", ""),
    ("text", "Third item after an empty paragraph"), ("space", " "),
    ("text", "Fourth item after a whitespace-only paragraph"), ("empty", ""), ("empty", ""),
]):
    par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    bullet(par)
    if kind != "empty":
        r = par.add_run()
        r.text = text
        r.font.size = Pt(24)
    else:
        end = OxmlElement("a:endParaRPr")
        end.set("sz", "2400")
        par._p.append(end)
p.save(sys.argv[1])
