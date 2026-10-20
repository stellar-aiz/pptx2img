"""Independent fixture for vertical text bodies (a:bodyPr@vert).

python-pptx is needed only for generation.

Every paragraph is left-aligned along its line.

- Slide 1: `vert270` labels in tall, narrow boxes with top, centre, and bottom
  anchors, and one `vert` label.
- Slide 2: `eaVert` bodies mixing Japanese and Latin text that is longer than one
  column, with top, centre, and bottom anchors, and one body with punctuation
  and brackets.
"""
import argparse
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt


def box(slide, x, y, w, h, vert, anchor, text, size, east_asian=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    style = shape._element.find(qn("p:style"))
    if style is not None:
        shape._element.remove(style)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(235, 242, 255)
    shape.line.color.rgb = RGBColor(40, 80, 160)
    shape.line.width = Pt(1)
    frame = shape.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Pt(4)
    frame.margin_top = frame.margin_bottom = Pt(4)
    body = frame._txBody.bodyPr
    body.set("vert", vert)
    body.set("anchor", anchor)
    # Start lines at the text's own left edge (the top of an eaVert column).
    frame.paragraphs[0].alignment = PP_ALIGN.LEFT
    run = frame.paragraphs[0].add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.name = "Arial"
    run.font.color.rgb = RGBColor(0, 0, 0)
    if east_asian:
        rpr = run._r.get_or_add_rPr()
        ea = rpr.makeelement(qn("a:ea"), {"typeface": "Yu Gothic"})
        latin = rpr.find(qn("a:latin"))
        latin.addnext(ea)
    return shape


def make(output):
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(8), Inches(4.5)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    for i, anchor in enumerate(["t", "ctr", "b"]):
        box(slide, 0.5 + i * 1.2, 0.5, 0.55, 3.2, "vert270", anchor, "Section " + "ABC"[i], 18)
    box(slide, 4.4, 0.5, 0.55, 3.2, "vert", "t", "Clockwise label", 18)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    text = "縦書きの本文はABCを含み、列の高さを超えて折り返します"
    for i, anchor in enumerate(["t", "ctr", "b"]):
        box(slide, 0.4 + i * 1.9, 0.4, 1.5, 2.6, "eaVert", anchor, text, 18, east_asian=True)
    box(slide, 6.1, 0.4, 1.5, 2.6, "eaVert", "t", "「括弧」と句読点、長音ー。", 18, east_asian=True)
    deck.save(output)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("output", type=Path)
    make(parser.parse_args().output)
