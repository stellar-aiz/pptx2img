"""Independent fixture: SVG pictures stored the way Office stores inserted icons.

One slide with three pictures of the same generated icon (an SVG written here):

1. SVG only: `a:blip` has no `r:embed`; the SVG is referenced from
   `a:blip/a:extLst/a:ext/asvg:svgBlip/@r:embed`;
2. SVG with a PNG fallback: `a:blip/@r:embed` points at a PNG drawn here with Pillow
   (a coarser copy of the icon) and `svgBlip` points at the SVG;
3. SVG only, stretched to a wide frame (the picture fills its frame, ignoring aspect).

rhodopsin renders the SVG for all three. python-pptx and Pillow are needed only for
generation.
"""
from io import BytesIO
from pathlib import Path
import sys

from lxml import etree
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.util import Inches, Pt

A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
ASVG = 'http://schemas.microsoft.com/office/drawing/2016/SVG/main'
SVG_EXT = '{96DAC541-7B7A-43D3-8B79-37D633B846F1}'

ICON = b'''<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 96 96" width="96" height="96">
<circle cx="48" cy="48" r="44" fill="#1A7F8E"/>
<path d="M26 50 L42 66 L72 32" fill="none" stroke="#FFFFFF" stroke-width="10" stroke-linecap="round" stroke-linejoin="round"/>
<rect x="8" y="84" width="80" height="8" fill="#E8A33D"/>
</svg>'''


def fallback_png():
    scale = 2
    image = Image.new('RGBA', (96 * scale, 96 * scale), (0, 0, 0, 0))
    draw = ImageDraw.Draw(image)
    draw.ellipse([4 * scale, 4 * scale, 92 * scale, 92 * scale], fill=(26, 127, 142))
    draw.line([(26 * scale, 50 * scale), (42 * scale, 66 * scale), (72 * scale, 32 * scale)],
              fill=(255, 255, 255), width=10 * scale, joint='curve')
    draw.rectangle([8 * scale, 84 * scale, 88 * scale, 92 * scale], fill=(232, 163, 61))
    data = BytesIO()
    image.save(data, 'PNG')
    data.seek(0)
    return data


def add_svg(slide, pic, svg_rid, keep_png):
    blip = pic._element.blipFill.blip
    if not keep_png:
        del blip.attrib[f'{{{R}}}embed']
    ext_list = etree.SubElement(blip, f'{{{A}}}extLst')
    ext = etree.SubElement(ext_list, f'{{{A}}}ext', uri=SVG_EXT)
    svg_blip = etree.SubElement(ext, f'{{{ASVG}}}svgBlip', nsmap={'asvg': ASVG})
    svg_blip.set(f'{{{R}}}embed', svg_rid)


def main(out):
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(10), Inches(5.625)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    part = Part(PackURI('/ppt/media/icon1.svg'), 'image/svg+xml', deck.part.package, ICON)
    svg_rid = slide.part.relate_to(part, RT.IMAGE)
    cases = [
        ('SVG only', Inches(0.6), Inches(2), Inches(2), False),
        ('SVG + PNG fallback', Inches(3.4), Inches(2), Inches(2), True),
        ('SVG only, stretched', Inches(6.2), Inches(3.2), Inches(1.6), False),
    ]
    for caption, left, width, height, keep_png in cases:
        pic = slide.shapes.add_picture(fallback_png(), left, Inches(1.0), width, height)
        add_svg(slide, pic, svg_rid, keep_png)
        box = slide.shapes.add_textbox(left, Inches(3.4), Inches(3), Inches(0.6))
        box.text_frame.text = caption
        box.text_frame.paragraphs[0].runs[0].font.size = Pt(16)
    deck.save(out)


if __name__ == '__main__':
    main(Path(sys.argv[1]))
