"""Independent fixtures comparing punctuation compression and inline positions; python-pptx is needed only for generation."""
import argparse
import json
from pathlib import Path
from pptx import Presentation
from pptx.util import Pt, Inches
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.oxml.xmlchemy import OxmlElement


def make(output, sizes, nested=False):
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(4), Inches(2)
    cases = []
    texts = ['「「あ」」い'] if nested else ['あ、いう、え', 'あ。いう。え', 'あ「い」う、']
    patterns = [(size, text, width) for size in sizes for text in texts
                for width in [size * 4 - .1, size * 4, size * 4 + .1, size * 4.5,
                              *[size * 5 + delta for delta in (-.1, -.07, -.05, -.04, -.03, -.02, -.01, 0, .01, .02, .05, .1)],
                              size * 6, size * 6 + 1]]
    for size, text, width in patterns:
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        shape = slide.shapes.add_textbox(Pt(24), Pt(15), Pt(width), Pt(110))
        frame = shape.text_frame
        frame.word_wrap, frame.auto_size = True, MSO_AUTO_SIZE.NONE
        frame.margin_left = frame.margin_right = 0
        frame.margin_top = frame.margin_bottom = 0
        para = frame.paragraphs[0]
        para._p.get_or_add_pPr().set('hangingPunct', '0')
        run = para.add_run()
        run.text = text
        run.font.name, run.font.size = 'Hiragino Sans', Pt(size)
        props = run._r.get_or_add_rPr()
        props.set('lang', 'ja-JP')
        east = OxmlElement('a:ea')
        east.set('typeface', 'Hiragino Sans')
        props.append(east)
        cases.append({'slide': len(cases) + 1, 'font_pt': size,
                      'text': text, 'width_pt': width})
    deck.save(output)
    output.with_suffix('.json').write_text(json.dumps(cases, ensure_ascii=False, indent=2) + '\n')


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('output', type=Path)
    parser.add_argument('--sizes', type=float, nargs='+', default=[12, 18, 24])
    parser.add_argument('--nested', action='store_true', help='Generate the known unresolved nested-bracket fixture')
    args = parser.parse_args()
    make(args.output, args.sizes, args.nested)
