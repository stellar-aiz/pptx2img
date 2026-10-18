"""テーマ線参照、同等の直接指定、細線の端部を比較する自作資料。"""
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZipFile
import sys
from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn
from pptx.util import Inches

A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
LINE = '<a:ln xmlns:a="'+A+'" w="25400" cap="flat"><a:solidFill><a:schemeClr val="phClr"/></a:solidFill><a:prstDash val="dash"/><a:round/><a:headEnd type="diamond"/><a:tailEnd type="triangle"/></a:ln>'

def main(output):
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(10), Inches(5.625)
    for mode in ['theme', 'explicit', 'thin']:
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        for index, y in enumerate([1, 2.7, 4.4]):
            shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(1), Inches(y), Inches(9), Inches(y))
            props = shape._element.spPr
            if props.find(qn('a:ln')) is not None:
                props.remove(props.find(qn('a:ln')))
            style = shape._element.find(qn('p:style'))
            for child in list(style):
                style.remove(child)
            reference = OxmlElement('a:lnRef')
            reference.set('idx', '1')
            color = OxmlElement('a:srgbClr')
            color.set('val', ['3366CC', '000000', 'B03030'][index])
            reference.append(color)
            style.append(reference)
            if mode == 'explicit':
                props.append(etree.fromstring(LINE.replace('<a:schemeClr val="phClr"/>', etree.tostring(color).decode())))
            elif mode == 'thin':
                line = OxmlElement('a:ln')
                line.set('w', '3175')
                dash = OxmlElement('a:prstDash'); dash.set('val', 'solid'); line.append(dash)
                head = OxmlElement('a:headEnd'); head.set('type', 'none'); line.append(head)
                tail = OxmlElement('a:tailEnd'); tail.set('type', 'triangle'); line.append(tail)
                props.append(line)
    with TemporaryDirectory() as folder:
        original = Path(folder)/'original.pptx'
        deck.save(original)
        with ZipFile(original) as source, ZipFile(output, 'w') as target:
            for entry in source.infolist():
                data = source.read(entry.filename)
                if entry.filename == 'ppt/theme/theme1.xml':
                    theme = etree.fromstring(data)
                    lines = theme.find('.//'+qn('a:lnStyleLst'))
                    for old in list(lines): lines.remove(old)
                    for _ in range(3): lines.append(etree.fromstring(LINE))
                    data = etree.tostring(theme, xml_declaration=True, encoding='UTF-8', standalone=True)
                target.writestr(entry, data)

if __name__ == '__main__':
    main(Path(sys.argv[1]))
